from datetime import timedelta

from conftest import NOW, RecordingNotifier
from openpyxl import load_workbook

from lead_hub.jobs import check_sla, notify_new_lead, send_digest
from lead_hub.models import LeadStatus
from lead_hub.schemas import LeadIn
from lead_hub.services import intake, update_lead

DAY = timedelta(hours=24)


def new_lead(session, phone="+393331234567", **kwargs):
    data = {"name": "Mario <b>", "phone": phone, "message": "Хочу узнать цену", "source": "landing"}
    data.update(kwargs)
    outcome = intake(session, LeadIn(**data), NOW, DAY)
    session.flush()
    return outcome.lead


def test_new_lead_goes_to_the_assigned_manager(session, managers, settings, notifier):
    lead = new_lead(session)
    assert notify_new_lead(lead, notifier, settings)
    chat, text = notifier.messages[0]
    assert chat == "111" and "Новая заявка" in text and "Ann" in text
    assert "Mario &lt;b&gt;" in text  # user input is HTML-escaped


def test_unassigned_lead_goes_to_the_admin(session, settings, notifier):
    notify_new_lead(new_lead(session), notifier, settings)
    assert notifier.messages[0][0] == "999"


def test_sla_reminder_once_then_escalation_once(session, managers, settings, notifier):
    lead = new_lead(session)
    assert check_sla(session, notifier, settings, NOW + timedelta(minutes=10)) == (0, 0)  # still within SLA

    assert check_sla(session, notifier, settings, NOW + timedelta(minutes=31)) == (1, 0)
    assert notifier.messages[-1][0] == "111" and "ждёт уже 31 мин" in notifier.messages[-1][1]
    assert check_sla(session, notifier, settings, NOW + timedelta(minutes=40)) == (0, 0)  # no spam

    assert check_sla(session, notifier, settings, NOW + timedelta(minutes=61)) == (0, 1)
    assert notifier.messages[-1][0] == "999" and "🚨" in notifier.messages[-1][1]
    assert check_sla(session, notifier, settings, NOW + timedelta(minutes=120)) == (0, 0)
    assert [e.kind for e in lead.events][-2:] == ["reminder", "escalation"]


def test_taken_leads_are_not_reminded(session, managers, settings, notifier):
    lead = new_lead(session)
    update_lead(session, lead, status=LeadStatus.IN_PROGRESS)
    session.flush()
    assert check_sla(session, notifier, settings, NOW + timedelta(hours=5)) == (0, 0)


def test_failed_send_is_retried_next_time(session, managers, settings):
    new_lead(session)
    broken = RecordingNotifier(fail=True)
    assert check_sla(session, broken, settings, NOW + timedelta(minutes=31)) == (0, 0)
    working = RecordingNotifier()
    assert check_sla(session, working, settings, NOW + timedelta(minutes=32)) == (1, 0)


def test_daily_digest_with_excel(session, managers, settings, notifier, tmp_path):
    new_lead(session, phone="+393330000001", message='=HYPERLINK("http://evil")')
    won = new_lead(session, phone="+393330000002")
    update_lead(session, won, status=LeadStatus.IN_PROGRESS)
    update_lead(session, won, status=LeadStatus.WON)
    session.flush()

    text = send_digest(session, notifier, settings, NOW + timedelta(hours=2), tmp_path)
    assert "Всего: 2" in text and "Успех: 1" in text and "landing: 2" in text
    chat, path, caption = notifier.documents[0]
    assert chat == "999" and caption == text
    ws = load_workbook(path)["Заявки"]
    assert ws.max_row == 3  # header + 2 leads
    assert ws["J3"].data_type == "s" or ws["J2"].data_type == "s"  # formula stays text
