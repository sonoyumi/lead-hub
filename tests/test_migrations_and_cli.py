from alembic.autogenerate import compare_metadata
from alembic.migration import MigrationContext
from conftest import make_settings
from openpyxl import load_workbook
from sqlalchemy import inspect

from lead_hub.cli import run, upgrade_database
from lead_hub.db import make_engine
from lead_hub.models import Base


def test_migrations_match_the_models(tmp_path):
    settings = make_settings(database_url=f"sqlite:///{tmp_path / 'm.db'}")
    upgrade_database(settings)
    engine = make_engine(settings.database_url)
    with engine.connect() as conn:
        assert set(inspect(conn).get_table_names()) >= {"managers", "leads", "lead_events", "alembic_version"}
        diff = compare_metadata(MigrationContext.configure(conn), Base.metadata)
    engine.dispose()
    assert diff == [], f"models and migrations differ: {diff}"


def test_cli_managers_and_export(tmp_path, capsys):
    settings = make_settings(database_url=f"sqlite:///{tmp_path / 'c.db'}")
    upgrade_database(settings)

    assert run(["managers", "add", "Ann", "--chat", "111"], settings) == 0
    assert run(["managers", "add", "Bob"], settings) == 0
    assert run(["managers", "disable", "2"], settings) == 0
    capsys.readouterr()
    run(["managers", "list"], settings)
    out = capsys.readouterr().out
    assert "#1  Ann  chat=111  активен" in out and "#2  Bob  chat=—  отключён" in out
    assert run(["managers", "disable", "99"], settings) == 2

    output = tmp_path / "leads.xlsx"
    assert run(["export", "--days", "7", "-o", str(output)], settings) == 0
    assert load_workbook(output)["Заявки"].max_row == 1  # header only: no leads yet

    assert run(["sla"], settings) == 0
    assert "Напоминаний: 0" in capsys.readouterr().out
