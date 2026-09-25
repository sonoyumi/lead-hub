"""Excel export of leads."""

from __future__ import annotations

from datetime import tzinfo
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from lead_hub.db import as_utc
from lead_hub.models import Lead

HEADERS = ["ID", "Создана", "Имя", "Телефон", "Email", "Источник", "Статус", "Менеджер", "Повторов", "Сообщение"]
STATUS_RU = {"new": "новая", "in_progress": "в работе", "won": "успех", "lost": "отказ"}


def export_leads(leads: list[Lead], path: Path, tz: tzinfo) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for lead in leads:
        created = as_utc(lead.created_at).astimezone(tz).replace(tzinfo=None)  # Excel has no time zones
        ws.append(
            [
                lead.id,
                created,
                lead.name,
                lead.phone or "",
                lead.email or "",
                lead.source,
                STATUS_RU.get(lead.status.value, lead.status.value),
                lead.manager.name if lead.manager else "",
                lead.duplicates,
                lead.message,
            ]
        )
        for cell in ws[ws.max_row]:
            # Text from web forms must never become a live formula (formula injection).
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
        ws.cell(ws.max_row, 2).number_format = "dd.mm.yyyy hh:mm"
    ws.freeze_panes = "A2"
    widths = [6, 17, 22, 16, 26, 14, 11, 18, 9, 60]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
